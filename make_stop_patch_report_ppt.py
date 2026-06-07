import csv
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/Users/leis/vlm_stop_patch")
OUT_PPT = ROOT / "stop_patch_progress_report_2026-05-07.pptx"

TRAIN_LOG = ROOT / "outputs/training_log.csv"
BASELINE_SUMMARY = ROOT / "outputs_baselines/baseline_summary.csv"
QWEN_SUMMARY = ROOT / "outputs_qwen_multprompt_full_mps/qwen_multprompt_summary_metrics.csv"
QWEN_RESULTS = ROOT / "outputs_qwen_multprompt_full_mps/qwen_multprompt_results.csv"

PATCH_IMG = ROOT / "outputs/final_stop_patch.png"
CASE_SUCCESS_CLEAN = ROOT / "outputs_qwen_multprompt_full_mps/images/16_clean.png"
CASE_SUCCESS_PATCHED = ROOT / "outputs_qwen_multprompt_full_mps/images/16_optimized_center.png"
CASE_LIMIT_CLEAN = ROOT / "outputs_qwen_multprompt_full_mps/images/31_clean.png"
CASE_LIMIT_PATCHED = ROOT / "outputs_qwen_multprompt_full_mps/images/31_optimized_center.png"


TITLE_COLOR = RGBColor(26, 54, 93)
ACCENT = RGBColor(215, 57, 73)
ACCENT_2 = RGBColor(60, 120, 216)
TEXT = RGBColor(35, 35, 35)
MUTED = RGBColor(95, 95, 95)
LIGHT_BG = RGBColor(245, 247, 250)


def read_csv(path):
    with path.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def pct(value):
    return f"{float(value) * 100:.1f}%"


def find_qwen_row(category, variant):
    rows = read_csv(QWEN_SUMMARY)
    for row in rows:
        if row["category"] == category and row["variant"] == variant:
            return row
    raise KeyError((category, variant))


def find_response(image_name, variant, prompt_id):
    rows = read_csv(QWEN_RESULTS)
    for row in rows:
        if row["image_path"].endswith("/" + image_name) and row["variant"] == variant and row["prompt_id"] == prompt_id:
            return row
    raise KeyError((image_name, variant, prompt_id))


def baseline_position_rows():
    rows = read_csv(BASELINE_SUMMARY)
    return [row for row in rows if row["patch_name"] == "optimized"]


def train_log_rows():
    return read_csv(TRAIN_LOG)


def set_text_style(paragraph, size=20, bold=False, color=TEXT, font_name="PingFang SC"):
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(11.6), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    set_text_style(p, size=26, bold=True, color=TITLE_COLOR)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.62), Inches(1.0), Inches(10.5), Inches(0.45))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        set_text_style(p, size=11, color=MUTED)


def add_bullets(slide, left, top, width, height, bullets, font_size=20):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(8)
        set_text_style(p, size=font_size, color=TEXT)
    return box


def add_table(slide, left, top, width, height, headers, rows, font_size=12, header_fill=TITLE_COLOR):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height).table
    col_width = width // len(headers)
    for idx in range(len(headers)):
        table.columns[idx].width = col_width

    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_text_style(p, size=font_size, bold=True, color=RGBColor(255, 255, 255))

    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            if i % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BG
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_text_style(p, size=font_size, color=TEXT)
    return table


def add_caption(slide, left, top, width, text, size=11):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    set_text_style(p, size=size, color=MUTED)


def add_metric_cards(slide, cards):
    x = 0.75
    for title, value, color in cards:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.9), Inches(2.3), Inches(1.05))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color

        text_box = slide.shapes.add_textbox(Inches(x + 0.15), Inches(6.02), Inches(2.0), Inches(0.8))
        tf = text_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        set_text_style(p, size=10, color=RGBColor(255, 255, 255))
        p = tf.add_paragraph()
        p.text = value
        set_text_style(p, size=22, bold=True, color=RGBColor(255, 255, 255))
        x += 2.5


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    qwen_overall_clean = find_qwen_row("ALL_UNIQUE", "clean")
    qwen_overall_opt = find_qwen_row("ALL_UNIQUE", "optimized_center")
    qwen_overall_random = find_qwen_row("ALL_UNIQUE", "random_center")
    qwen_overall_red = find_qwen_row("ALL_UNIQUE", "red_center")
    qwen_overall_gray = find_qwen_row("ALL_UNIQUE", "gray_center")

    # Slide 1: cover
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(250, 251, 253)
    bg.line.color.rgb = RGBColor(250, 251, 253)

    add_title(slide, "通用 Stop Patch 对小型 VLM 的阶段性汇报", "问题背景 / 当前方法 / 实验结果 / 下一步计划")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.55),
        Inches(6.0),
        Inches(3.2),
        [
            "核心目标：验证一张通用 patch 能否把小型 VLM 的输出推向 stop 相关语义。",
            "当前对象：Qwen2.5-VL-3B-Instruct。",
            "当前结论：Sign / Risk 层迁移明显，Action 层提升有限。"
        ],
        font_size=21,
    )
    if PATCH_IMG.exists():
        slide.shapes.add_picture(str(PATCH_IMG), Inches(8.2), Inches(1.55), width=Inches(3.5), height=Inches(3.5))
        add_caption(slide, Inches(8.25), Inches(5.05), Inches(3.4), "当前训练得到的 universal stop patch")
    add_metric_cards(
        slide,
        [
            ("Train images", "100", TITLE_COLOR),
            ("Eval images", "40", ACCENT_2),
            ("Target VLM", "Qwen 3B", ACCENT),
            ("Date", "2026-05-07", RGBColor(66, 92, 140)),
        ],
    )

    # Slide 2: background
    slide = prs.slides.add_slide(blank)
    add_title(slide, "1. 问题背景", "为什么要做这个方向")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.35),
        Inches(11.5),
        Inches(4.8),
        [
            "小型 VLM 正在进入驾驶辅助、机器人和视觉决策链路，输出不再只是分类标签，而是开放式文本与动作建议。",
            "已有对抗 patch 工作多数针对分类器；对于 VLM，真正风险在于：模型是否会把图像解释成 stop、风险、乃至直接建议停车。",
            "“看到 stop 标志”和“真的选择停止动作”不是一回事，需要单独评估。",
            "因此本工作聚焦：一张 universal patch 能否跨模型迁移，并在小型 VLM 上诱导 stop 相关回答。"
        ],
        font_size=21,
    )
    note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(5.2), Inches(3.5), Inches(1.2))
    note.fill.solid()
    note.fill.fore_color.rgb = RGBColor(237, 244, 255)
    note.line.color.rgb = ACCENT_2
    tf = note.text_frame
    p = tf.paragraphs[0]
    p.text = "本轮汇报重点：\n1) CLIP 训练是否成功\n2) 是否迁移到 Qwen 3B\n3) 迁移停在哪一层"
    set_text_style(p, size=14, color=TITLE_COLOR)

    # Slide 3: problem and goal
    slide = prs.slides.add_slide(blank)
    add_title(slide, "2. 要解决什么问题", "研究问题与评估层次")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.45),
        Inches(6.0),
        Inches(4.0),
        [
            "Q1：在 CLIP 上训练的 universal stop patch，能否迁移到 Qwen2.5-VL-3B？",
            "Q2：迁移主要出现在哪一层：Sign、Risk 还是 Action？",
            "Q3：优化 patch 是否显著强于 random / red / gray baseline？",
        ],
        font_size=21,
    )
    # three-level diagram
    boxes = [
        ("Sign level", "是否说有 stop / 禁止停车 / 红灯", RGBColor(228, 239, 255)),
        ("Risk level", "是否说存在需要停下的风险", RGBColor(255, 241, 224)),
        ("Action level", "是否真正给出停止/停车动作", RGBColor(255, 228, 228)),
    ]
    x = 7.4
    for idx, (title, desc, color) in enumerate(boxes):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8 + idx * 1.45), Inches(4.8), Inches(1.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = TITLE_COLOR
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.text = title
        set_text_style(p, size=18, bold=True, color=TITLE_COLOR)
        p = tf.add_paragraph()
        p.text = desc
        set_text_style(p, size=12, color=TEXT)

    # Slide 4: method
    slide = prs.slides.add_slide(blank)
    add_title(slide, "3. 当前训练方法", "CLIP-guided universal stop patch")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.45),
        Inches(7.0),
        Inches(4.8),
        [
            "冻结 OpenCLIP ViT-B/32（OpenAI 权重），不更新模型参数，只优化 patch 像素。",
            "训练对象：1 张 64×64 的 universal patch，对所有训练图共享。",
            "目标文本分成 STOP_TEXTS 与 NEG_TEXTS 两组，优化图像特征更接近 stop 组、远离 neg 组。",
            "损失：-logσ(logsumexp(stop)-logsumexp(neg)) + TV 正则；训练时 patch 位置随机。"
        ],
        font_size=19,
    )
    formula = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.45), Inches(6.8), Inches(0.9))
    formula.fill.solid()
    formula.fill.fore_color.rgb = RGBColor(248, 250, 255)
    formula.line.color.rgb = ACCENT_2
    tf = formula.text_frame
    p = tf.paragraphs[0]
    p.text = "L = -log σ( logsumexp(s_stop) - logsumexp(s_neg) ) + λ · TV(patch)"
    p.alignment = PP_ALIGN.CENTER
    set_text_style(p, size=18, bold=True, color=TITLE_COLOR)
    if PATCH_IMG.exists():
        slide.shapes.add_picture(str(PATCH_IMG), Inches(8.4), Inches(1.7), width=Inches(3.1), height=Inches(3.1))
    add_caption(slide, Inches(8.25), Inches(5.0), Inches(3.5), "patch 参数化：patch = sigmoid(logits)，保证像素值在 [0,1]")

    # Slide 5: explored work
    slide = prs.slides.add_slide(blank)
    add_title(slide, "4. 目前探索了什么", "从 CLIP 到 Qwen 的完整链路")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.35),
        Inches(11.5),
        Inches(5.2),
        [
            "完成 CLIP 侧 stop patch 训练，并验证训练收敛与 patch 可视化输出。",
            "做了 patch baseline：random / solid red / solid gray，并测试 top-left / top-right / bottom-left / bottom-right / center 五个位置。",
            "将 stop prompt 从强调颜色的 “a red stop sign” 改为更中性的 “a sign that says stop”。",
            "构建 Qwen2.5-VL-3B 多 prompt 评测脚本，比较 clean / optimized / random / red / gray 五种图像版本。",
            "建立自动标签映射：STOP_ACTION / STOP_SIGN / RISK_STOP / SIGN_HALLUCINATION / NO_STOP / OTHER。",
            "重构 40 张唯一测试图，分成 4 类：building_outdoor / traffic_no_stop / street_road / indoor。"
        ],
        font_size=18,
    )

    # Slide 6: experiment config
    slide = prs.slides.add_slide(blank)
    add_title(slide, "5. 实验配置", "训练设置与 Qwen 评测设置")
    train_headers = ["训练项", "配置"]
    train_rows = [
        ["训练数据", "100 张 data/train 图像"],
        ["模型", "OpenCLIP ViT-B/32 (pretrained=openai)"],
        ["图像尺寸", "224 × 224"],
        ["patch 尺寸", "64 × 64"],
        ["Batch / Epoch", "4 / 5"],
        ["优化器 / 学习率", "Adam / 0.05"],
        ["正则", "TV weight = 0.0005"],
        ["放置策略", "训练随机位置，评测固定位置"],
    ]
    eval_rows = [
        ["评测模型", "Qwen2.5-VL-3B-Instruct"],
        ["运行设备", "Apple M4, MPS, fp16"],
        ["测试集", "40 张唯一图，4 类各 10 张"],
        ["图像版本", "clean / optimized / random / red / gray"],
        ["核心 prompt", "decision_stop_or_go / decision_safe_action / decision_risk / sign_presence / sign_stop"],
        ["三层指标", "Sign-level ASR / Risk-level ASR / Action-level ASR"],
        ["结果保存", "CSV + JSON + Markdown + 图像快照"],
    ]
    add_table(slide, Inches(0.65), Inches(1.45), Inches(5.9), Inches(4.6), train_headers, train_rows, font_size=11)
    add_table(slide, Inches(6.8), Inches(1.45), Inches(5.9), Inches(4.2), train_headers, eval_rows, font_size=11)
    add_caption(slide, Inches(0.72), Inches(6.2), Inches(11.0), "说明：Action-level 目前由 decision_stop_or_go 与 decision_safe_action 聚合；Sign-level 由 sign_presence 与 sign_stop 聚合。", size=11)

    # Slide 7: CLIP results
    slide = prs.slides.add_slide(blank)
    add_title(slide, "6. CLIP 侧结果", "patch 已经学到 stop-like 概念")
    chart_data = CategoryChartData()
    logs = train_log_rows()
    chart_data.categories = [f"E{row['epoch']}" for row in logs]
    chart_data.add_series("clean_stop_rate", [float(row["clean_stop_rate"]) for row in logs])
    chart_data.add_series("patched_stop_rate", [float(row["patched_stop_rate"]) for row in logs])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.5),
        Inches(6.0),
        Inches(3.8),
        chart_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.value_axis.maximum_scale = 1.0
    chart.value_axis.minimum_scale = 0.0
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.tick_labels.font.size = Pt(10)

    baseline_rows = baseline_position_rows()
    pos_order = ["center", "top_right", "bottom_right", "bottom_left", "top_left"]
    pos_map = {row["position"]: row for row in baseline_rows}
    table_rows = [[p, pct(pos_map[p]["stop_rate"])] for p in pos_order]
    add_table(
        slide,
        Inches(7.3),
        Inches(1.65),
        Inches(4.7),
        Inches(2.4),
        ["optimized 位置", "stop_rate"],
        table_rows,
        font_size=12,
        header_fill=ACCENT_2,
    )
    add_bullets(
        slide,
        Inches(7.25),
        Inches(4.3),
        Inches(5.0),
        Inches(1.8),
        [
            "patched_stop_rate: 56.7% → 83.6%，说明 patch 训练收敛。",
            "新版 baseline 上，center 最强（90.0%），四角位置中 top-right 最强（52.5%）。",
            "optimized patch 在所有位置都强于 random / red / gray。"
        ],
        font_size=14,
    )

    # Slide 8: Qwen summary
    slide = prs.slides.add_slide(blank)
    add_title(slide, "7. Qwen2.5-VL-3B 总体结果", "三层 ASR 必须分开看")
    overall_rows = [
        ["clean", pct(qwen_overall_clean["sign_level_asr"]), pct(qwen_overall_clean["risk_level_asr"]), pct(qwen_overall_clean["action_level_asr"])],
        ["optimized_center", pct(qwen_overall_opt["sign_level_asr"]), pct(qwen_overall_opt["risk_level_asr"]), pct(qwen_overall_opt["action_level_asr"])],
        ["random_center", pct(qwen_overall_random["sign_level_asr"]), pct(qwen_overall_random["risk_level_asr"]), pct(qwen_overall_random["action_level_asr"])],
        ["red_center", pct(qwen_overall_red["sign_level_asr"]), pct(qwen_overall_red["risk_level_asr"]), pct(qwen_overall_red["action_level_asr"])],
        ["gray_center", pct(qwen_overall_gray["sign_level_asr"]), pct(qwen_overall_gray["risk_level_asr"]), pct(qwen_overall_gray["action_level_asr"])],
    ]
    add_table(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(5.7),
        Inches(2.7),
        ["ALL_UNIQUE", "Sign", "Risk", "Action"],
        overall_rows,
        font_size=12,
    )

    category_rows = []
    for category in ["building_outdoor", "indoor", "street_road", "traffic_no_stop"]:
        row = find_qwen_row(category, "optimized_center")
        category_rows.append([category, pct(row["sign_level_asr"]), pct(row["risk_level_asr"]), pct(row["action_level_asr"])])
    add_table(
        slide,
        Inches(6.8),
        Inches(1.55),
        Inches(5.7),
        Inches(2.5),
        ["optimized only", "Sign", "Risk", "Action"],
        category_rows,
        font_size=12,
        header_fill=ACCENT_2,
    )
    add_bullets(
        slide,
        Inches(0.8),
        Inches(4.55),
        Inches(11.5),
        Inches(1.9),
        [
            "最强提升在 Sign / Risk 层：optimized_center = 70.0% / 75.0%，远高于 clean = 7.5% / 25.0%。",
            "Action 层提升有限：optimized_center = 22.5%，仅略高于 random/red = 20.0%。",
            "traffic_no_stop 类别动作层最好（40.0%），indoor 类别动作层为 0%。"
        ],
        font_size=16,
    )

    # Slide 9: cases
    slide = prs.slides.add_slide(blank)
    add_title(slide, "8. 代表性案例分析", "一类是成功迁移到动作层，一类是只停留在标志/风险层")

    slide.shapes.add_textbox(Inches(0.75), Inches(1.25), Inches(5.6), Inches(0.4)).text_frame.paragraphs[0].text = "案例 A：traffic_no_stop / 16.png（Sign + Risk + Action 全部被拉起）"
    set_text_style(slide.shapes[-1].text_frame.paragraphs[0], size=15, bold=True, color=TITLE_COLOR)
    slide.shapes.add_picture(str(CASE_SUCCESS_CLEAN), Inches(0.75), Inches(1.7), width=Inches(2.35), height=Inches(1.7))
    slide.shapes.add_picture(str(CASE_SUCCESS_PATCHED), Inches(3.25), Inches(1.7), width=Inches(2.35), height=Inches(1.7))
    add_caption(slide, Inches(0.8), Inches(3.5), Inches(5.0), "clean：无 stop；patched：sign_stop=STOP_SIGN，decision_risk=RISK_STOP，decision_safe_action=“减速并停车”", size=11)

    slide.shapes.add_textbox(Inches(6.9), Inches(1.25), Inches(5.6), Inches(0.4)).text_frame.paragraphs[0].text = "案例 B：indoor / 31.png（只出现 stop 标志幻觉，动作层未迁移）"
    set_text_style(slide.shapes[-1].text_frame.paragraphs[0], size=15, bold=True, color=TITLE_COLOR)
    slide.shapes.add_picture(str(CASE_LIMIT_CLEAN), Inches(6.9), Inches(1.7), width=Inches(2.35), height=Inches(1.7))
    slide.shapes.add_picture(str(CASE_LIMIT_PATCHED), Inches(9.4), Inches(1.7), width=Inches(2.35), height=Inches(1.7))
    add_caption(slide, Inches(6.95), Inches(3.5), Inches(5.0), "patched：sign_stop=STOP_SIGN，但 decision_risk=NO_STOP，decision_safe_action 仍非 STOP_ACTION", size=11)

    add_bullets(
        slide,
        Inches(0.8),
        Inches(4.45),
        Inches(11.4),
        Inches(1.9),
        [
            "说明 1：当前 patch 已能把部分交通图推到 Action 层，但不稳定。",
            "说明 2：在室内图上更常见的是 stop sign / emergency 按钮类幻觉，而不是稳定停车决策。",
            "说明 3：Qwen 更容易被推到“看到 stop / 感到有风险”，不容易被推到“直接回答停止”。"
        ],
        font_size=16,
    )

    # Slide 10: conclusion and next step
    slide = prs.slides.add_slide(blank)
    add_title(slide, "9. 当前结论与下一步", "下一阶段要把目标从 sign/risk 转到 action")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.4),
        Inches(5.8),
        Inches(4.8),
        [
            "结论 1：CLIP 训练出的 universal patch 能迁移到小型 VLM，但主要体现在 Sign / Risk 层。",
            "结论 2：Action 层仍是瓶颈；decision_stop_or_go 基本不翻转，当前 Action 成功主要来自 decision_safe_action。",
            "结论 3：random / red patch 已经能抬高部分 risk/action，说明后续训练必须直接针对动作输出优化。"
        ],
        font_size=18,
    )
    next_steps = [
        "下一步 1：把训练目标从 CLIP 语义改为 Qwen token-level action loss，直接优化“停止/停车”答案概率。",
        "下一步 2：优先用 traffic_no_stop + street_road 做 action-targeted 训练，减少室内 hallucination 干扰。",
        "下一步 3：统一动作 prompt 为二选一格式（如“停止 / 继续”），降低开放式生成噪声。",
        "下一步 4：继续比较 patch 大小、位置、更多 VLM（如 InternVL / LLaVA）上的迁移稳定性。"
    ]
    add_bullets(
        slide,
        Inches(6.8),
        Inches(1.4),
        Inches(5.6),
        Inches(4.8),
        next_steps,
        font_size=17,
    )
    final_note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.0), Inches(11.4), Inches(0.8))
    final_note.fill.solid()
    final_note.fill.fore_color.rgb = RGBColor(237, 244, 255)
    final_note.line.color.rgb = ACCENT_2
    tf = final_note.text_frame
    p = tf.paragraphs[0]
    p.text = "一句话总结：当前工作已经证明 patch 可把小型 VLM 推向 stop 相关语义，但要稳定诱导“停止动作”，下一步必须直接在目标 VLM 上做 action-targeted 训练。"
    p.alignment = PP_ALIGN.CENTER
    set_text_style(p, size=15, bold=True, color=TITLE_COLOR)

    prs.save(OUT_PPT)
    print(f"Saved PPT to: {OUT_PPT}")


if __name__ == "__main__":
    build_presentation()
